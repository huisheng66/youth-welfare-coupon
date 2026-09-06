"""Bulk-import file parsing: turn xlsx / csv / txt / docx uploads into string rows.

All parsing helpers raise ``ValueError`` with a user-facing Chinese message;
endpoints convert that to HTTP 400.
"""

from __future__ import annotations

import csv
import io
import re

from sqlalchemy.orm import Session

from app.models.entities import Account, Role, UserProfile

KIND_USERS = "users"
KIND_ISSUE = "issue"
KIND_POINTS = "points"

MAX_UPLOAD_BYTES = 2 * 1024 * 1024
# 解析阶段资源上限（T14）：列数、单元格长度、Office 解压体积与压缩比
MAX_COLUMNS_PER_ROW = 64
MAX_CELL_LENGTH = 512
MAX_DECOMPRESSED_BYTES = 50 * 1024 * 1024
MAX_COMPRESSION_RATIO = 100

# 无表头时的固定列序（每种导入类型）
KIND_POSITIONAL_COLUMNS: dict[str, tuple[str, ...]] = {
    KIND_USERS: ("real_name", "student_no", "username", "phone", "organization", "remark"),
    KIND_ISSUE: ("identifier",),
    KIND_POINTS: ("identifier", "hours", "reason"),
}

# 表头关键词 → 逻辑列名（匹配为「关键词包含于单元格文本」；顺序即优先级）
HEADER_KEYWORDS: dict[str, tuple[str, ...]] = {
    "real_name": ("姓名", "名字", "真实姓名", "real_name"),
    "student_no": ("学号", "student_no", "student"),
    "username": ("用户名", "账号", "登录名", "username"),
    "phone": ("手机", "电话", "phone", "mobile"),
    "email": ("邮箱", "email", "e-mail"),
    "organization": ("组织", "单位", "机构", "organization"),
    "remark": ("备注", "remark", "note"),
    "identifier": ("用户标识", "标识", "用户", "identifier"),
    "hours": ("时长", "小时", "工时", "hours"),
    "reason": ("说明", "事由", "reason"),
}

# 单元格内容分隔符（txt / docx 段落的一行文本拆多列）
_LINE_SPLIT_RE = re.compile(r"\t|,|，")


def read_upload_bytes(spooled, *, max_bytes: int = MAX_UPLOAD_BYTES) -> bytes:
    """Read an uploaded file with a hard size cap. Raises ValueError on oversize/empty."""
    data = spooled.read(max_bytes + 1)
    if len(data) > max_bytes:
        raise ValueError(f"文件过大（上限 {max_bytes // 1024} KB）")
    if not data:
        raise ValueError("文件为空")
    return data


def decode_text(data: bytes) -> str:
    """utf-8(-sig) 优先，回退 GBK（中文 Excel 另存的 CSV 常见编码）。"""
    for enc in ("utf-8-sig", "utf-8", "gbk"):
        try:
            return data.decode(enc)
        except UnicodeDecodeError:
            continue
    raise ValueError("文件编码无法识别（请用 UTF-8 或 GBK 保存）")


def _cell_str(v) -> str:
    """Normalize a spreadsheet cell to text; keep integer-valued floats clean."""
    if v is None:
        return ""
    if isinstance(v, float) and v.is_integer():
        return str(int(v))
    return str(v).strip()


def _check_zip_sizes(data: bytes) -> None:
    """Office 文件（zip）解压体积与压缩比检查，快速拒绝 zip 炸弹。

    在 openpyxl/python-docx 解析前执行：仅读 zip 目录，不解压内容。
    """
    import zipfile

    try:
        with zipfile.ZipFile(io.BytesIO(data)) as zf:
            total = sum(info.file_size for info in zf.infolist())
            compressed = sum(info.compress_size for info in zf.infolist())
    except zipfile.BadZipFile:
        return  # 非 zip 内容交给后续解析器报「无法解析」
    if total > MAX_DECOMPRESSED_BYTES:
        raise ValueError(
            f"文件解压后体积过大（上限 {MAX_DECOMPRESSED_BYTES // 1024 // 1024} MB），已拒绝解析"
        )
    if compressed > 0 and total / max(compressed, 1) > MAX_COMPRESSION_RATIO:
        raise ValueError("文件压缩比异常（疑似压缩炸弹），已拒绝解析")


def _parse_xlsx(data: bytes) -> list[list[str]]:
    from openpyxl import load_workbook

    _check_zip_sizes(data)
    try:
        wb = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
        try:
            ws = wb.worksheets[0]
            return [[_cell_str(v) for v in row] for row in ws.iter_rows(values_only=True)]
        finally:
            wb.close()
    except ValueError:
        raise
    except Exception as exc:  # BadZipFile / KeyError / InvalidFileException 等 → 统一 400
        raise ValueError(f"xlsx 文件无法解析：{exc}") from exc


def _parse_docx(data: bytes) -> list[list[str]]:
    import docx  # python-docx

    _check_zip_sizes(data)
    try:
        document = docx.Document(io.BytesIO(data))
    except ValueError:
        raise
    except Exception as exc:  # BadZipFile / PackageNotFoundError 等 → 统一 400
        raise ValueError(f"docx 文件无法解析：{exc}") from exc
    if document.tables:
        return [[_cell_str(cell.text) for cell in row.cells] for row in document.tables[0].rows]
    # 无表格：每个非空段落是一行，含分隔符时拆多列
    rows: list[list[str]] = []
    for para in document.paragraphs:
        text = (para.text or "").strip()
        if not text:
            continue
        rows.append([part.strip() for part in _LINE_SPLIT_RE.split(text)])
    return rows


def _parse_csv(data: bytes) -> list[list[str]]:
    return [[_cell_str(c) for c in row] for row in csv.reader(io.StringIO(decode_text(data)))]


def _parse_txt(data: bytes) -> list[list[str]]:
    rows: list[list[str]] = []
    for line in decode_text(data).splitlines():
        line = line.strip()
        if not line:
            continue
        rows.append([part.strip() for part in _LINE_SPLIT_RE.split(line)])
    return rows


def parse_table_file(filename: str, data: bytes, *, max_rows: int) -> list[list[str]]:
    """Parse an uploaded table file into stripped non-empty rows of strings."""
    ext = filename.rsplit(".", 1)[-1].lower() if "." in filename else ""
    if ext == "xlsx":
        rows = _parse_xlsx(data)
    elif ext == "docx":
        rows = _parse_docx(data)
    elif ext == "csv":
        rows = _parse_csv(data)
    elif ext in {"txt", "text"}:
        rows = _parse_txt(data)
    else:
        raise ValueError("不支持的文件格式（支持 .xlsx / .csv / .txt / .docx）")

    cleaned = [row for row in rows if any(row)]
    if not cleaned:
        raise ValueError("文件中没有数据")
    if len(cleaned) > max_rows:
        raise ValueError(f"数据行数 {len(cleaned)} 超过上限 {max_rows}，请分批导入")
    for lineno, row in enumerate(cleaned, start=1):
        if len(row) > MAX_COLUMNS_PER_ROW:
            raise ValueError(f"第 {lineno} 行列数 {len(row)} 超过上限 {MAX_COLUMNS_PER_ROW}")
        for cell in row:
            if len(cell) > MAX_CELL_LENGTH:
                raise ValueError(f"第 {lineno} 行存在超长单元格（上限 {MAX_CELL_LENGTH} 字符）")
    return cleaned


def map_columns(first_row: list[str], kind: str) -> dict[str, int] | None:
    """Detect a header row and map logical names → column indexes.

    Returns None when the first row holds data (positional order applies).
    For issue/points kinds an explicit 用户名/手机/学号 header doubles as
    the identifier column.
    """

    def hit(cell: str, keywords: tuple[str, ...]) -> bool:
        lc = cell.lower()
        return any(kw.lower() in lc for kw in keywords)

    if not any(hit(c, kws) for c in first_row for kws in HEADER_KEYWORDS.values()):
        return None

    mapping: dict[str, int] = {}
    for idx, cell in enumerate(first_row):
        for logical, kws in HEADER_KEYWORDS.items():
            if logical in mapping:
                continue
            if hit(cell, kws):
                mapping[logical] = idx
                break

    if kind in {KIND_ISSUE, KIND_POINTS} and "identifier" not in mapping:
        for alt in ("username", "phone", "email", "student_no"):
            if alt in mapping:
                mapping["identifier"] = mapping[alt]
                break
    if kind == KIND_POINTS and "reason" not in mapping and "remark" in mapping:
        mapping["reason"] = mapping["remark"]
    return mapping


def resolve_user(db: Session, token: str) -> Account | None:
    """按 用户名 → 邮箱 → 手机 → 学号 精确匹配 role=user 账号。

    学号可能重复（转学/重号），匹配不唯一或查无此人时返回 None，
    由调用方记为行级错误。
    """
    account, _reason = resolve_user_detailed(db, token)
    return account


def resolve_user_detailed(db: Session, token: str) -> tuple[Account | None, str | None]:
    """标识解析（T14）：返回 (账号, 失败原因)。

    匹配优先级：用户名 → 邮箱 → 手机 → 学号（姓名不作为唯一标识）。
    - 同一标识跨字段命中不同用户 → 歧义，必须报告而不是静默取第一个；
    - 学号命中多个用户 → 歧义；
    - 全部未命中 → 查无此人。
    """
    token = (token or "").strip()
    if not token:
        return None, "用户标识为空"
    q = db.query(Account).filter(Account.role == Role.user)
    hits: dict[str, Account] = {}

    account = q.filter(Account.username == token).first()
    if account:
        hits[account.id] = account
    account = q.filter(Account.email == token.lower()).first()
    if account:
        hits[account.id] = account
    account = q.filter(Account.phone == token).first()
    if account:
        hits[account.id] = account
    if not hits:
        profiles = (
            db.query(UserProfile)
            .join(Account, UserProfile.account_id == Account.id)
            .filter(UserProfile.student_no == token, Account.role == Role.user)
            .limit(2)
            .all()
        )
        if len(profiles) > 1:
            return None, "学号匹配到多个用户，存在歧义，请改用用户名/手机/邮箱"
        if len(profiles) == 1:
            account = db.get(Account, profiles[0].account_id)
            if account:
                hits[account.id] = account

    if len(hits) > 1:
        return None, "该标识在不同字段命中多个用户，存在歧义，请改用更精确的标识"
    if hits:
        return next(iter(hits.values())), None
    return None, "用户不存在或无法唯一识别（支持用户名/邮箱/手机/学号）"
