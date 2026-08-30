"""批量导入：用户名单 / 按名单发券 / 按名单时长（xlsx / csv / txt / docx）。"""

from __future__ import annotations

import unittest
from decimal import Decimal
from io import BytesIO

from tests._helpers import TempApp, fresh_settings, reset_env_defaults


def _csv(text: str) -> bytes:
    return text.strip().encode("utf-8")


def _xlsx(rows: list[list]) -> bytes:
    from openpyxl import Workbook

    wb = Workbook()
    for row in rows:
        wb.active.append(row)
    buf = BytesIO()
    wb.save(buf)
    return buf.getvalue()


def _docx(rows: list[list[str]]) -> bytes:
    import docx as docx_lib

    d = docx_lib.Document()
    table = d.add_table(rows=len(rows), cols=len(rows[0]))
    for i, row in enumerate(rows):
        for j, cell in enumerate(row):
            table.cell(i, j).text = cell
    buf = BytesIO()
    d.save(buf)
    return buf.getvalue()


def _student_no_of(ta: TempApp, username: str) -> str:
    with ta.session() as db:
        from app.models.entities import Account, UserProfile

        acc = db.query(Account).filter(Account.username == username).first()
        profile = db.query(UserProfile).filter(UserProfile.account_id == acc.id).first()
        return profile.student_no


def _user_id_of(ta: TempApp, username: str) -> str:
    with ta.session() as db:
        from app.models.entities import Account

        return db.query(Account).filter(Account.username == username).first().id


def _coupon_count(ta: TempApp, username: str) -> int:
    with ta.session() as db:
        from app.models.entities import Account, CouponInstance

        acc = db.query(Account).filter(Account.username == username).first()
        return db.query(CouponInstance).filter(CouponInstance.user_id == acc.id).count()


def _active_template_id(ta: TempApp, token: str) -> str:
    with ta.client() as c:
        r = c.get(
            "/api/coupons/templates",
            headers=ta.bearer(token),
            params={"active_only": True},
        )
    assert r.status_code == 200, r.text
    return r.json()["items"][0]["id"]


def _post_import(ta: TempApp, token: str, path: str, filename: str, data: bytes, form: dict | None = None):
    with ta.client() as c:
        return c.post(
            f"/api{path}",
            headers=ta.bearer(token),
            files={"file": (filename, data, "application/octet-stream")},
            data=form or {},
        )


class TestUserListImport(unittest.TestCase):
    def tearDown(self) -> None:
        reset_env_defaults()

    def test_import_csv_with_header(self) -> None:
        """带表头 CSV：导入即已通过、初始密码可登录、出现在已通过列表。"""
        data = _csv(
            "姓名,学号,用户名,手机,组织,备注\n"
            "张三,20260001,zhangsan,13800000001,某某大学,\n"
            "李四,20260002,lisi,,某某大学,班长\n"
        )
        with TempApp() as ta:
            token = ta.login("issuer", "issuer123")
            r = _post_import(ta, token, "/users/import", "users.csv", data, {"dry_run": "false"})
            self.assertEqual(r.status_code, 200, r.text)
            body = r.json()
            self.assertEqual(body["total"], 2)
            self.assertEqual(body["succeeded"], 2)
            self.assertEqual(body["failed"], 0)
            self.assertEqual(body["default_password"], "youth123456")

            with ta.client() as c:
                lst = c.get(
                    "/api/users",
                    headers=ta.bearer(token),
                    params={"verify_status": "approved", "limit": 100},
                )
                items = {u["username"]: u for u in lst.json()["items"]}
                self.assertIn("zhangsan", items)
                self.assertIn("lisi", items)
                self.assertEqual(items["zhangsan"]["real_name"], "张三")
                self.assertEqual(items["lisi"]["student_no"], "20260002")

                # 导入用户可立即用初始密码登录，且须强制改密
                login = c.post("/api/auth/login", json={"username": "zhangsan", "password": "youth123456"})
                self.assertEqual(login.status_code, 200, login.text)
                me = c.get("/api/auth/me", headers={"Authorization": f"Bearer {login.json()['access_token']}"})
                self.assertEqual(me.status_code, 200, me.text)
                self.assertTrue(me.json()["must_change_password"], me.text)

    def test_import_txt_tab_without_header(self) -> None:
        """无表头 txt（tab 分隔）：用户名缺省回退学号。"""
        data = "张三\t20260001\t\t13800000001\t某某大学\t\n李四\t20260002\t\t\t\t".encode("utf-8")
        with TempApp() as ta:
            token = ta.login("admin", "admin123")
            r = _post_import(ta, token, "/users/import", "users.txt", data)
            self.assertEqual(r.status_code, 200, r.text)
            body = r.json()
            self.assertEqual(body["succeeded"], 2, body)
            with ta.client() as c:
                lst = c.get("/api/users", headers=ta.bearer(token), params={"q": "20260002"})
                usernames = [u["username"] for u in lst.json()["items"]]
                self.assertIn("20260002", usernames)  # 用户名回退为学号

    def test_import_xlsx_docx_and_gbk_csv(self) -> None:
        """xlsx / docx 表格、GBK 编码 CSV 三种格式均可解析导入。"""
        with TempApp() as ta:
            token = ta.login("issuer", "issuer123")
            r = _post_import(
                ta, token, "/users/import", "users.xlsx",
                _xlsx([["姓名", "学号", "用户名", "手机"], ["王五", "20260003", "wangwu", 13800000003]]),
            )
            self.assertEqual(r.status_code, 200, r.text)
            self.assertEqual(r.json()["succeeded"], 1, r.text)

        with TempApp() as ta:
            token = ta.login("issuer", "issuer123")
            r = _post_import(
                ta, token, "/users/import", "users.docx",
                _docx([["姓名", "学号", "用户名"], ["赵六", "20260004", "zhaoliu"]]),
            )
            self.assertEqual(r.status_code, 200, r.text)
            self.assertEqual(r.json()["succeeded"], 1, r.text)

        with TempApp() as ta:
            token = ta.login("issuer", "issuer123")
            gbk = "姓名,学号,用户名,手机\n钱七,20260005,qianqi,13800000005\n".encode("gbk")
            r = _post_import(ta, token, "/users/import", "users.csv", gbk)
            self.assertEqual(r.status_code, 200, r.text)
            self.assertEqual(r.json()["succeeded"], 1, r.text)

    def test_import_rejects_duplicates(self) -> None:
        """库内用户名冲突与文件内手机重复均记行错误。"""
        data = _csv(
            "姓名,学号,用户名,手机\n"
            "张三,20260001,lisi,13800000001\n"
            "李四,20260002,zhangsan2,13800000001\n"
            "王五,20260003,youth1,13800000002\n"
        )
        with TempApp() as ta:
            token = ta.login("admin", "admin123")
            r = _post_import(ta, token, "/users/import", "users.csv", data)
            self.assertEqual(r.status_code, 200, r.text)
            body = r.json()
            self.assertEqual(body["succeeded"], 1, body)
            reasons = [e["reason"] for e in body["errors"]]
            self.assertIn("手机号在文件内重复", reasons)
            self.assertIn("用户名已存在", reasons)
            for err in body["errors"]:
                self.assertTrue(err["row"] >= 1 and err["identifier"])

    def test_import_dry_run_creates_nothing(self) -> None:
        data = _csv("姓名,学号,用户名\n张三,20260001,zhangsan\n")
        with TempApp() as ta:
            token = ta.login("admin", "admin123")
            r = _post_import(ta, token, "/users/import", "users.csv", data, {"dry_run": "true"})
            self.assertEqual(r.status_code, 200, r.text)
            self.assertEqual(r.json()["succeeded"], 1, r.text)
            with ta.session() as db:
                from app.models.entities import Account

                self.assertIsNone(db.query(Account).filter(Account.username == "zhangsan").first())
            with ta.client() as c:
                login = c.post("/api/auth/login", json={"username": "zhangsan", "password": "youth123456"})
                self.assertEqual(login.status_code, 400, login.text)

    def test_import_skips_email_notify_without_smtp(self) -> None:
        """控制台模式（未配 SMTP）：不排队邮件，结果中说明跳过。"""
        data = _csv("姓名,用户名,邮箱\n张三,zhangsan,zhangsan@example.com\n")
        with TempApp() as ta:
            token = ta.login("admin", "admin123")
            r = _post_import(ta, token, "/users/import", "users.csv", data)
            self.assertEqual(r.status_code, 200, r.text)
            body = r.json()
            self.assertEqual(body["succeeded"], 1, body)
            self.assertEqual(body["email_queued"], 0)
            self.assertIn("跳过邮件通知", body["message"])

    def test_import_sends_password_emails_when_smtp_configured(self) -> None:
        """配置 SMTP 后：含邮箱行在响应后由后台任务发送开通邮件（含初始密码）。"""
        import asyncio
        from unittest import mock

        data = _csv("姓名,用户名,邮箱\n张三,zhangsan,zhangsan@example.com\n李四,lisi,\n")
        with TempApp() as ta:
            fresh_settings(
                MAIL_SERVER="smtp.example.invalid",
                MAIL_PORT="465",
                MAIL_USERNAME="noreply@example.invalid",
                MAIL_PASSWORD="app-password",
                MAIL_FROM="noreply@example.invalid",
            )
            try:
                token = ta.login("admin", "admin123")
                with mock.patch("aiosmtplib.send", new_callable=mock.AsyncMock) as smtp_send, ta.client() as c:
                    r = c.post(
                        "/api/users/import",
                        headers=ta.bearer(token),
                        files={"file": ("users.csv", data, "text/csv")},
                        data={"dry_run": "false", "notify": "true"},
                    )
                    self.assertEqual(r.status_code, 200, r.text)
                    body = r.json()
                    self.assertEqual(body["succeeded"], 2, body)
                    self.assertEqual(body["email_queued"], 1)
                    self.assertIn("已排队向 1 人发送开通邮件", body["message"])

                    # TestClient 会等后台任务完成：仅含邮箱的那行收到开通邮件
                    smtp_send.assert_awaited_once()
                    message = smtp_send.await_args.args[0]
                    kwargs = smtp_send.await_args.kwargs
                    self.assertEqual(kwargs["recipients"], ["zhangsan@example.com"])
                    import base64

                    html_part = message.get_payload()[1].get_payload()
                    html = base64.b64decode(html_part).decode("utf-8")  # 含中文，传输为 base64
                    self.assertIn("zhangsan", html)
                    self.assertIn("youth123456", html)
                    self.assertIn("修改密码", html)
            finally:
                fresh_settings(
                    MAIL_SERVER=None,
                    MAIL_USERNAME=None,
                    MAIL_PASSWORD=None,
                    MAIL_FROM=None,
                )

    def test_import_rejects_bad_files(self) -> None:
        with TempApp() as ta:
            token = ta.login("admin", "admin123")
            # 不支持的扩展名
            r = _post_import(ta, token, "/users/import", "list.pdf", b"%PDF-1.4")
            self.assertEqual(r.status_code, 400, r.text)
            self.assertIn("不支持的文件格式", r.json()["detail"])
            # 空文件
            r = _post_import(ta, token, "/users/import", "users.csv", b"")
            self.assertEqual(r.status_code, 400, r.text)
            # 超过 2MB 上限
            r = _post_import(ta, token, "/users/import", "users.txt", b"a" * (2 * 1024 * 1024 + 1))
            self.assertEqual(r.status_code, 400, r.text)
            self.assertIn("文件过大", r.json()["detail"])

    def test_import_rejects_corrupt_office_files(self) -> None:
        """伪装 / 损坏的 xlsx、docx → 400 而非 500。"""
        import io
        import zipfile

        buf = io.BytesIO()
        with zipfile.ZipFile(buf, "w") as z:
            z.writestr("xl/workbook.xml", "broken")
        broken_zip = buf.getvalue()
        with TempApp() as ta:
            token = ta.login("admin", "admin123")
            for filename, data in (
                ("evil.xlsx", b"not a real zip"),
                ("broken.xlsx", broken_zip),
                ("evil.docx", b"not a real zip"),
            ):
                r = _post_import(ta, token, "/users/import", filename, data)
                self.assertEqual(r.status_code, 400, f"{filename}: {r.status_code} {r.text[:120]}")
                self.assertIn("无法解析", r.json()["detail"])

    def test_import_row_limit_enforced(self) -> None:
        fresh_settings(IMPORT_MAX_ROWS="2")
        try:
            data = _csv("姓名,用户名\n张三,a001\n李四,a002\n王五,a003\n")
            with TempApp() as ta:
                token = ta.login("admin", "admin123")
                r = _post_import(ta, token, "/users/import", "users.csv", data)
                self.assertEqual(r.status_code, 400, r.text)
                self.assertIn("超过上限", r.json()["detail"])
        finally:
            fresh_settings(IMPORT_MAX_ROWS=None)

    def test_import_forbidden_for_lower_roles(self) -> None:
        data = _csv("姓名,用户名\n张三,a001\n")
        with TempApp() as ta:
            merchant_token = ta.login("merchant1", "merchant123")
            youth_token = ta.login("youth1", "youth123")
            r = _post_import(ta, merchant_token, "/users/import", "users.csv", data)
            self.assertEqual(r.status_code, 403, r.text)
            r = _post_import(ta, youth_token, "/users/import", "users.csv", data)
            self.assertEqual(r.status_code, 403, r.text)
            with ta.client() as c:  # 未登录
                r = c.post(
                    "/api/users/import",
                    files={"file": ("users.csv", data, "text/csv")},
                )
                self.assertEqual(r.status_code, 401, r.text)


class TestIssueByListImport(unittest.TestCase):
    def tearDown(self) -> None:
        reset_env_defaults()

    def test_issue_import_mixed_rows(self) -> None:
        """学号可定位用户；待核验与未知用户进 failed，不中断。"""
        with TempApp() as ta:
            token = ta.login("issuer", "issuer123")
            template_id = _active_template_id(ta, token)
            student_no = _student_no_of(ta, "youth1")
            before = _coupon_count(ta, "youth1")  # seed 会预置 1 张演示券，按增量断言
            data = _csv(f"用户标识\n{student_no}\nyouth2\nnosuchuser\n")
            r = _post_import(
                ta, token, "/coupons/issue-import", "issue.csv", data,
                {"template_id": template_id, "quantity": "2"},
            )
            self.assertEqual(r.status_code, 200, r.text)
            body = r.json()
            self.assertEqual(body["total"], 3)
            self.assertEqual(body["succeeded"], 1, body)
            self.assertEqual(body["failed"], 2)
            reasons = {e["identifier"]: e["reason"] for e in body["errors"]}
            self.assertIn("仅可为已核验通过的用户发券", reasons["youth2"])
            self.assertIn("用户不存在", reasons["nosuchuser"])

            self.assertEqual(_coupon_count(ta, "youth1") - before, 2)
            self.assertIn("2 张券", body["message"])

    def test_issue_import_bad_template_rejected(self) -> None:
        with TempApp() as ta:
            token = ta.login("issuer", "issuer123")
            r = _post_import(
                ta, token, "/coupons/issue-import", "issue.csv", _csv("用户标识\nyouth1\n"),
                {"template_id": "not-exist", "quantity": "1"},
            )
            self.assertEqual(r.status_code, 400, r.text)
            self.assertIn("券模板不可用", r.json()["detail"])


class TestPointsByListImport(unittest.TestCase):
    def tearDown(self) -> None:
        reset_env_defaults()

    def test_grant_import_mixed_rows(self) -> None:
        """正常入账、待核验失败、格式错误与零时长均逐行处理。"""
        with TempApp() as ta:
            token = ta.login("admin", "admin123")
            youth1_id = _user_id_of(ta, "youth1")
            data = _csv(
                "用户标识,时长(小时),说明\n"
                "youth1,2.5,社区志愿服务\n"
                "youth2,1,\n"
                "youth1,abc,\n"
                "youth1,0,\n"
            )
            r = _post_import(ta, token, "/points/grant-import", "points.csv", data, {"reason": "批量导入"})
            self.assertEqual(r.status_code, 200, r.text)
            body = r.json()
            self.assertEqual(body["total"], 4)
            self.assertEqual(body["succeeded"], 1, body)
            self.assertEqual(body["failed"], 3)
            reasons = [e["reason"] for e in body["errors"]]
            self.assertTrue(any("仅可为已核验用户" in x for x in reasons), reasons)
            self.assertTrue(any("时长格式无效" in x for x in reasons), reasons)
            self.assertTrue(any("不能为 0" in x for x in reasons), reasons)

            with ta.client() as c:
                bal = c.get(f"/api/points/users/{youth1_id}", headers=ta.bearer(token))
                self.assertEqual(bal.status_code, 200, bal.text)
                self.assertEqual(Decimal(bal.json()["balance"]), Decimal("12.50"))
                # 流水中记录了导入的说明
                ledger = c.get("/api/points/ledger", headers=ta.bearer(token), params={"limit": 50})
                reasons_in_ledger = [i["reason"] for i in ledger.json()["items"]]
                self.assertIn("社区志愿服务", reasons_in_ledger)

    def test_grant_import_negative_overflow_fails_per_row(self) -> None:
        """负数扣减超出余额 → 行错误，不影响其他行。"""
        with TempApp() as ta:
            token = ta.login("admin", "admin123")
            data = _csv("用户标识,时长(小时)\nyouth1,-999\n")
            r = _post_import(ta, token, "/points/grant-import", "points.csv", data)
            self.assertEqual(r.status_code, 200, r.text)
            body = r.json()
            self.assertEqual(body["succeeded"], 0, body)
            self.assertTrue(any("余额不足" in e["reason"] for e in body["errors"]), body)


if __name__ == "__main__":
    unittest.main()
